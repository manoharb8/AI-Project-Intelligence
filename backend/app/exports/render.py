"""Consistent, paginated PDF and editable Word renderers for the shared report."""
from io import BytesIO
from pathlib import Path
from threading import Lock
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, LongTable, TableStyle, Image, KeepTogether
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from app.exports.report import Report
from app.exports.matrix import matrix_png

FONT_DIR=Path(__file__).with_name('fonts');FONT_LOCK=Lock()
NAVY='182C42';TEAL='176B75';GOLD='B58A35';INK='253341';MUTED='526372'


def column_weights(headers):
    if headers==['Risk ID and title','Probability / Impact','Severity / Status','Owner / Target date']:return [31,22,20,27]
    if headers and headers[-1]=='Evidence':return [45,19,23,13] if len(headers)==4 else [100/len(headers)]*len(headers)
    if headers and headers[0]=='Dimension':return [23,40,12,25]
    if len(headers)==3:return [25,35,40]
    return [100/len(headers)]*len(headers)


def word_bytes(report: Report)->bytes:
    document=Document();section=document.sections[0]
    section.page_width,section.page_height=Inches(8.27),Inches(11.69)
    section.top_margin=section.bottom_margin=Inches(.72);section.left_margin=section.right_margin=Inches(.72)
    section.header_distance=section.footer_distance=Inches(.3)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Quote','Caption'):
        style=document.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor.from_string(NAVY if name.startswith('Heading') or name=='Title' else INK)
    for style in document.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    normal=document.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
    for name,size in [('Title',24),('Subtitle',12),('Heading 1',16),('Heading 2',12),('Heading 3',10.5),('Caption',9),('Quote',10)]:
        style=document.styles[name];style.font.size=Pt(size)
        style.paragraph_format.keep_with_next=name in ('Title','Subtitle','Heading 1','Heading 2','Heading 3')
        if name.startswith('Heading'):style.paragraph_format.space_before=Pt(14);style.paragraph_format.space_after=Pt(7)
    document.styles['Subtitle'].font.color.rgb=RGBColor.from_string(TEAL)
    document.styles['Quote'].font.italic=False;document.styles['Quote'].paragraph_format.left_indent=Inches(.1)
    document.styles['Caption'].font.italic=False;document.styles['Caption'].font.color.rgb=RGBColor.from_string(MUTED)
    mapping={'title':'Title','subtitle':'Subtitle','heading':'Heading 1','subheading':'Heading 2','item':'Heading 3','quote':'Quote','caption':'Caption','note':'Caption'}
    for block in report.blocks:
        if block.kind=='table':
            table=document.add_table(rows=1,cols=len(block.headers));table.autofit=False
            weights=column_weights(block.headers)
            for column,w in zip(table.columns,weights):column.width=Inches(6.83*w/100)
            repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row_index,values in enumerate([block.headers,*block.rows]):
                cells=table.rows[0].cells if row_index==0 else table.add_row().cells
                for cell,value,w in zip(cells,values,weights):
                    cell.width=Inches(6.83*w/100);cell.text=value
                    tcpr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),NAVY if row_index==0 else ('F2F6F7' if row_index%2 else 'FFFFFF'));tcpr.append(shade)
                    borders=OxmlElement('w:tcBorders')
                    for side in ('top','left','bottom','right'):
                        edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
                    tcpr.append(borders)
                    margins=OxmlElement('w:tcMar')
                    for side in ('top','left','bottom','right'):
                        margin=OxmlElement('w:'+side);margin.set(qn('w:w'),'100');margin.set(qn('w:type'),'dxa');margins.append(margin)
                    tcpr.append(margins)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.08
                        for run in p.runs:run.font.size=Pt(9);run.bold=row_index==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if row_index==0 else INK)
            document.add_paragraph().paragraph_format.space_after=Pt(2)
        elif block.kind=='chart':
            p=document.add_paragraph();p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(BytesIO(matrix_png(block.payload)),width=Inches(6.8))
            document.add_paragraph('Risk positions from stated source ratings. Full values and references follow in the register.','Caption')
        else:
            p=document.add_paragraph(block.text,mapping.get(block.kind,'Normal'));p.paragraph_format.widow_control=True
    header=section.header.paragraphs[0];header.text='PROJECT INTELLIGENCE  /  EVIDENCE REPORT'
    for run in header.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string(TEAL)
    footer=section.footer.paragraphs[0];footer.add_run('Project evidence report  |  Page ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    for run in footer.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string(MUTED)
    document.core_properties.title=f'{report.workspace} - Project evidence report';document.core_properties.author='Project Intelligence & Risk Advisor'
    buffer=BytesIO();document.save(buffer);return buffer.getvalue()


def pdf_bytes(report: Report)->bytes:
    with FONT_LOCK:
        if 'ProjectReport' not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont('ProjectReport',str(FONT_DIR/'DejaVuSans.ttf')));pdfmetrics.registerFont(TTFont('ProjectReportBold',str(FONT_DIR/'DejaVuSans-Bold.ttf')))
    glyphs=pdfmetrics.getFont('ProjectReport').face.charToGlyph
    for b in report.blocks:
        texts=[b.text,*b.headers,*[v for row in b.rows for v in row]]
        if b.kind=='chart':texts += [r.source_id for r,_,_ in b.payload[1]]
        if any(ord(c) not in glyphs for text in texts for c in text if not c.isspace()):raise ValueError('PDF cannot render one or more source characters with the bundled font. Download Word to preserve the original text.')
    buffer=BytesIO();width=A4[0]-104
    pdf=SimpleDocTemplate(buffer,pagesize=A4,rightMargin=52,leftMargin=52,topMargin=48,bottomMargin=48,title=f'{report.workspace} - Project evidence report',author='Project Intelligence & Risk Advisor')
    normal=ParagraphStyle('Body',fontName='ProjectReport',fontSize=10,leading=14.5,spaceAfter=7,textColor=colors.HexColor('#'+INK),splitLongWords=True)
    styles={'body':normal,'note':ParagraphStyle('Note',parent=normal,fontSize=9,leading=12,textColor=colors.HexColor('#'+MUTED)),
        'caption':ParagraphStyle('Caption',parent=normal,fontSize=8.8,leading=12,textColor=colors.HexColor('#'+MUTED)),
        'quote':ParagraphStyle('Quote',parent=normal,fontSize=9.3,leading=13.2,leftIndent=5),
        'title':ParagraphStyle('Title',parent=normal,fontName='ProjectReportBold',fontSize=23,leading=28,spaceAfter=10,keepWithNext=True),
        'subtitle':ParagraphStyle('Subtitle',parent=normal,fontSize=12,leading=17,textColor=colors.HexColor('#'+TEAL),spaceAfter=14,keepWithNext=True)}
    for kind,size in [('heading',16),('subheading',12),('item',10.2)]:styles[kind]=ParagraphStyle(kind,parent=normal,fontName='ProjectReportBold',fontSize=size,leading=size+5,textColor=colors.HexColor('#'+NAVY),spaceBefore=17 if kind=='heading' else 9,spaceAfter=7,keepWithNext=True)
    cell=ParagraphStyle('Cell',parent=normal,fontSize=9,leading=12,spaceAfter=0)
    head=ParagraphStyle('Head',parent=cell,fontName='ProjectReportBold',textColor=colors.white)
    def paragraph(text,style):return Paragraph(escape(text).replace('\n','<br/>'),style)
    story=[]
    for b in report.blocks:
        if b.kind=='table':
            data=[[paragraph(v,head) for v in b.headers]]+[[paragraph(v,cell) for v in row] for row in b.rows]
            table=LongTable(data,colWidths=[width*w/100 for w in column_weights(b.headers)],repeatRows=1,hAlign='LEFT',splitByRow=1,splitInRow=1)
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#'+NAVY)),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#F2F6F7'),colors.white]),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]));table.spaceAfter=12;story.append(table)
        elif b.kind=='chart':
            story.append(Image(BytesIO(matrix_png(b.payload)),width=width,height=width*1370/2100))
            story.append(paragraph('Risk positions from stated source ratings. Full values and references follow in the register.',styles['caption']))
        else:story.append(paragraph(b.text,styles.get(b.kind,normal)))
    def decoration(canvas,doc):
        canvas.saveState();canvas.setFont('ProjectReport',7.5);canvas.setFillColor(colors.HexColor('#'+TEAL));canvas.drawString(52,A4[1]-25,'PROJECT INTELLIGENCE  /  EVIDENCE REPORT')
        canvas.setStrokeColor(colors.HexColor('#'+GOLD));canvas.setLineWidth(.7);canvas.line(52,36,A4[0]-52,36)
        canvas.setFillColor(colors.HexColor('#'+MUTED));canvas.drawString(52,23,'Project evidence report');canvas.drawRightString(A4[0]-52,23,f'Page {doc.page}');canvas.restoreState()
    pdf.build(story,onFirstPage=decoration,onLaterPages=decoration);return buffer.getvalue()
