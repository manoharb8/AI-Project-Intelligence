"""Real analysis -> saved workspace -> both report formats, without generation on export."""
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
import pytest
from docx import Document
from pypdf import PdfReader
from fastapi.testclient import TestClient
from app.main import create_app
from app.llm.provider import ExtractiveProvider
from app.services.embeddings import WordLlamaEmbeddings
from app.services.knowledge_base import KnowledgeBaseService
from app.exports.report import build_report, Report, Block
from app.exports.render import word_bytes, pdf_bytes

ROOT = Path(__file__).resolve().parents[2]

@pytest.fixture(scope='module')
def prepared(tmp_path_factory):
    kb = KnowledgeBaseService(tmp_path_factory.mktemp('exports'), WordLlamaEmbeddings())
    for path in sorted((ROOT / 'samples').iterdir()):
        kb.ingest(path.name, path.read_bytes())
    with TestClient(create_app(kb, ExtractiveProvider())) as client:
        for endpoint in ('agents/scope','agents/risk','agents/blockers','documentation/generate','project-health/assess'):
            response = client.post('/api/'+endpoint,json={})
            assert response.status_code == 200, response.text
        for question in ('Are we on track?','What is blocking the project?'):
            assert client.post('/api/assistant/query',json={'question':question}).status_code == 200
        yield client


def text_of(content, format):
    if format == 'pdf':
        return '\n'.join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
    return '\n'.join(p.text for p in Document(BytesIO(content)).paragraphs)

@pytest.mark.parametrize('format',['pdf','docx'])
def test_download_preserves_real_scores_unknowns_evidence_and_does_not_generate(prepared, format):
    manager = prepared.app.state.workspaces
    before = manager.snapshot('legacy').model_dump_json()
    snapshot = manager.snapshot('legacy')
    health = next(r.result for r in snapshot.results if r.kind=='health')
    report = build_report('Current project',snapshot)
    with patch.object(prepared.app.state.provider,'generate',side_effect=AssertionError('Export must not call provider')):
        response = prepared.get('/api/workspaces/legacy/export',params={'format':format},headers={'X-Workspace-ID':'legacy'})
    assert response.status_code == 200, response.text
    assert response.headers['cache-control']=='no-store'
    assert 'attachment;' in response.headers['content-disposition']
    text = text_of(response.content,format)
    # Ignore PDF line wrapping only; exact underlying content remains in report model.
    normal = ' '.join(text.split())
    for block in report.blocks:
        if block.text.startswith('Exported at (UTC):'):
            continue
        assert ' '.join(block.text.split()) in normal, block.text
    assert f'{health.overall_score} / 100' in text
    assert 'Not stated' in text
    assert 'Higher timeline risk = more risk' in text
    assert {e.document_type for r in snapshot.results for e in r.result.evidence} == {'pdf','docx','csv','txt'}
    assert manager.snapshot('legacy').model_dump_json()==before


def test_empty_workspace_cannot_export_another_projects_results(prepared):
    created=prepared.post('/api/workspaces',json={'name':'Empty export workspace'}).json()['id']
    response=prepared.get(f'/api/workspaces/{created}/export?format=pdf',headers={'X-Workspace-ID':created})
    assert response.status_code==409
    assert prepared.get(f'/api/workspaces/{created}/export?format=pdf',headers={'X-Workspace-ID':'legacy'}).status_code==409
    assert prepared.get('/api/workspaces/unknown/export?format=pdf').status_code==404
    assert prepared.get('/api/workspaces/legacy/export?format=html').status_code==422


def test_partial_and_stale_reports_label_missing_sections_without_fabricated_scores(prepared):
    snapshot=prepared.app.state.workspaces.snapshot('legacy').model_copy(deep=True)
    snapshot.results=[r for r in snapshot.results if r.kind=='scope']
    snapshot.source_version='changed'
    report=build_report('Partial report',snapshot)
    text='\n'.join(b.text for b in report.blocks)
    assert 'OUT OF DATE' in text
    assert 'No saved result is available' in text
    assert 'Overall health score: 0' not in text
    assert '6. Project health' in text


def test_insufficient_health_is_not_zero(prepared):
    snapshot=prepared.app.state.workspaces.snapshot('legacy').model_copy(deep=True)
    h=next(r for r in snapshot.results if r.kind=='health')
    h.result.overall_score=None
    h.result.dimensions[1].score=None
    text='\n'.join(b.text for b in build_report('Unknown',snapshot).blocks)
    assert 'Overall health score: Insufficient evidence' in text
    assert 'Score: Insufficient evidence' in text


def test_export_isolates_second_project_and_excludes_followup_chat(prepared):
    ident=prepared.post('/api/workspaces',json={'name':'Separate export project'}).json()['id']
    headers={'X-Workspace-ID':ident}
    prepared.post('/api/documents/upload',headers=headers,files={'files':('separate.txt',b'Project goal: Zephyr catalog. Deliverables: Zephyr index. Action item: review catalog.')})
    prepared.post('/api/agents/scope',headers=headers,json={})
    response=prepared.get(f'/api/workspaces/{ident}/export?format=docx',headers=headers)
    assert response.status_code==200
    text=text_of(response.content,'docx')
    assert 'Zephyr' in text and 'separate.txt' in text
    assert 'project_proposal.pdf' not in text and 'Ravi' not in text
    assert prepared.app.state.workspaces.snapshot('legacy').workspace_id=='legacy'

@pytest.mark.parametrize('format',['pdf','docx'])
def test_literal_markup_unicode_and_long_evidence_survive_rendering(format):
    original='José & Zoë <review> “API” — résumé.\n' + 'LongReferenceWithoutSpaces'*50
    report=Report('Rendering check','now',[Block('title','Rendering check'),Block('quote',original)])
    result=(pdf_bytes if format=='pdf' else word_bytes)(report)
    assert ''.join(original.split()) in ''.join(text_of(result,format).split())


def test_pdf_unsupported_glyph_is_explicit_not_silent_data_loss():
    report=Report('Unicode','now',[Block('body','项目资料')])
    with pytest.raises(ValueError,match='Download Word'):
        pdf_bytes(report)
    assert '项目资料' in text_of(word_bytes(report),'docx')


def test_report_excludes_optional_followup_conversation(prepared):
    manager=prepared.app.state.workspaces
    snapshot=manager.snapshot('legacy').model_copy(deep=True)
    entry=next(r for r in snapshot.results if r.kind=='assistant').model_copy(deep=True)
    entry.query='Who owns it?'
    entry.result.answer='Private follow-up content for exclusion test'
    snapshot.results.append(entry)
    text='\n'.join(b.text for b in build_report('Privacy check',snapshot).blocks)
    assert 'Private follow-up content' not in text


def test_processing_export_is_rejected_and_no_state_is_modified(prepared):
    manager=prepared.app.state.workspaces
    before=manager.snapshot('legacy').model_dump_json()
    with patch.object(manager.kb('legacy'),'list_documents',return_value=[{'status':'processing'}]):
        response=prepared.get('/api/workspaces/legacy/export?format=pdf')
    assert response.status_code==409
    assert 'still processing' in response.json()['detail']
    assert manager.snapshot('legacy').model_dump_json()==before


def test_word_title_has_no_decorative_rule():
    document=Document(BytesIO(word_bytes(Report('Style','now',[Block('title','Project report')]))))
    assert not list(document.styles['Title'].element.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr'))
