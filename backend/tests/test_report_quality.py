"""Export presentation correctness; synthetic edge cases are tests only."""
from io import BytesIO
from types import SimpleNamespace as S
import pytest
from docx import Document
from pypdf import PdfReader
from PIL import Image
from app.exports.risks import project_risks,axis,charts
from app.exports.matrix import matrix_png
from app.exports.report import Report,Block
from app.exports.render import pdf_bytes,word_bytes


def risk(id='R-01',date='2026-10-08',status='Open',likelihood='4/5 — High',impact='5/5 — Very High',severity='Critical'):
    quote=f'Risk ID: {id}; Risk Title: Workspace access; Owner: Reviewer; Target Date: 2026-10-10; As of: {date}.'
    return S(title='Workspace access',description=quote,likelihood=likelihood,impact=impact,severity=severity,status=status,mitigation='Review access',references=[S(evidence_id=date,quote=quote)])


def test_canonical_risks_keep_latest_snapshot_and_historical_evidence():
    records=project_risks([risk(date='2026-10-05'),risk(status='In Progress')])
    assert len(records)==1 and records[0].status=='In Progress'
    assert len(records[0].references)==2 and records[0].due=='2026-10-10'
    assert records[0].likelihood=='4/5 — High'
    assert records[0].impact=='5/5 — Very High'


def test_same_day_conflicting_values_remain_unknown_and_unpositioned():
    records=project_risks([risk(),risk(impact='2/5 — Low')])
    assert records[0].impact is None and 'impact' in records[0].conflicts
    groups,missing=charts(records);assert not groups and len(missing)==1


def test_axes_preserve_explicit_scales_without_guessing():
    assert axis('4/5 — High','probability')==('score5',.7)
    assert axis('5/5 — Very High','impact')==('score5',.9)
    assert axis('Medium','probability')==axis('Moderate','probability')
    assert axis('Likely','probability')==('words',.75)
    assert axis('100%','impact')==('percent',1)
    for value in ('4','6/5','High risk','Very serious','0.8','High or Low'):
        assert axis(value,'probability') is None


def test_coincident_nodes_are_all_rendered_and_input_values_unchanged():
    records=project_risks([risk(id=f'R-{i:02}') for i in range(1,17)])
    groups,missing=charts(records);assert len(groups)==1 and not missing
    image=Image.open(BytesIO(matrix_png(groups[0])))
    assert image.size==(2100,1370)
    assert all(x==.7 and y==.9 for _,x,y in groups[0][1])
    assert [r.source_id for r,_,_ in groups[0][1]]==[f'R-{i:02}' for i in range(1,17)]


def test_mixed_scales_separate_plots_and_missing_values_are_retained():
    records=project_risks([risk(),risk(id='R-02',likelihood='High',impact='High'),risk(id='R-03',likelihood=None)])
    groups,missing=charts(records)
    assert len(groups)==2 and [r.source_id for r in missing]==['R-03']


def test_both_renderers_have_same_table_values_and_word_has_repeated_header():
    report=Report('Test','now',[Block('title','Export check'),Block('table','',['Risk','Probability','Impact'],[['R-02','4/5 — High','5/5 — Very High'],['R-03','Not stated','High']])])
    word=Document(BytesIO(word_bytes(report)));pdf=PdfReader(BytesIO(pdf_bytes(report)))
    text=' '.join(p.extract_text() for p in pdf.pages)
    assert len(word.tables)==1
    assert word.tables[0].rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    for row in report.blocks[1].rows:
        for value in row:assert value in text and any(value in c.text for r in word.tables[0].rows for c in r.cells)


def test_graph_is_embedded_in_both_formats():
    group,_=charts(project_risks([risk()]))
    report=Report('Test','now',[Block('title','Risk matrix'),Block('chart','Probability × Impact',payload=group[0])])
    word=Document(BytesIO(word_bytes(report)));pdf=PdfReader(BytesIO(pdf_bytes(report)))
    assert len(word.inline_shapes)==1
    assert sum(len(page.images) for page in pdf.pages)==1
