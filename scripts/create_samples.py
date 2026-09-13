"""Generate controlled demonstration fixtures, not real project claims."""
from pathlib import Path
import csv
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'samples'
OUT.mkdir(exist_ok=True)

proposal = [
 'Atlas Portal - Project Proposal (controlled demo)',
 'The project goal is to build a web portal for customer support tickets.',
 'The portal will let customers submit tickets and track support responses.',
 'Milestones: requirements sign-off on 12 September 2026; API prototype on 20 September 2026;',
 'integration testing on 27 September 2026; pilot delivery on 30 September 2026.',
 'Deliverables: a React frontend, a ticket API, a database schema, and a testing report.',
 'Priya owns backend API development. Arjun owns the frontend. Meera owns testing.',
]
pdf=canvas.Canvas(str(OUT/'project_proposal.pdf'), pagesize=A4)
pdf.setTitle('Atlas Portal - Controlled Demo Proposal')
y=795
for i,line in enumerate(proposal):
 pdf.setFont('Helvetica-Bold' if i==0 else 'Helvetica',14 if i==0 else 10)
 pdf.drawString(45,y,line)
 y-=34 if i==0 else 24
pdf.save()

doc=Document()
doc.add_heading('Atlas Portal - Meeting notes',0)
doc.add_paragraph('Controlled demo. Meeting date: 15 September 2026.')
doc.add_paragraph('Current project blockers: the staging server is unavailable. Integration testing cannot start until access is restored.')
doc.add_paragraph('Pending decision: the product owner has not approved the ticket retention period. This remains unresolved.')
doc.add_paragraph('Action item: Ravi will restore staging server access by 18 September 2026. Status: pending.')
doc.add_paragraph('Action item: arrange the retention-policy review. Owner and due date were not recorded.')
table=doc.add_table(rows=1,cols=3)
for cell,text in zip(table.rows[0].cells,['Issue','Owner','Status']):cell.text=text
for cell,text in zip(table.add_row().cells,['SMTP credentials missing','Priya','Open']):cell.text=text
doc.save(OUT/'meeting_notes.docx')

with (OUT/'task_list.csv').open('w',newline='',encoding='utf-8') as f:
 writer=csv.writer(f)
 writer.writerow(['Task','Owner','Status','Due date','Dependency'])
 writer.writerow(['Build the ticket API','Priya','In progress','2026-09-20','Approved database schema'])
 writer.writerow(['Implement ticket form','Arjun','Incomplete','2026-09-22','Ticket API contract'])
 writer.writerow(['Run integration tests','Meera','Blocked','2026-09-27','Staging server access'])
 writer.writerow(['Approve requirements','Sana','Completed','2026-09-12','Stakeholder review'])

(OUT/'sprint_update.txt').write_text('''Atlas Portal - Sprint update (controlled demo)

Schedule risks: API development is three days behind plan because the database schema has not been approved. This delay threatens the integration testing milestone on 27 September 2026. No revised delivery date is confirmed.

Dependency gap: frontend integration depends on the ticket API contract. The contract is still pending review by Priya. The missing contract creates a delivery challenge for Arjun.

Completed progress: the login screen and the requirements review are finished. Incomplete tasks include the ticket API and ticket form. The staging server remains unavailable.
''',encoding='utf-8')
print('Created four controlled sample documents in',OUT)
