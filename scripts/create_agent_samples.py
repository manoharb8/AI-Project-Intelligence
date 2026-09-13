"""Four equivalent controlled fixtures for per-format Milestone 2 validation."""
from pathlib import Path
import csv
from docx import Document
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4
from xml.sax.saxutils import escape

STATEMENTS=[
 'Goal: build a web portal for customer support tickets.',
 'Milestone: the API prototype is scheduled for 20 September 2026.',
 'Timeline: integration testing is due 27 September 2026.',
 'Responsibility: Priya owns backend API development.',
 'Deliverables: a ticket API, a React frontend, and an integration test report.',
 'Schedule risk: API implementation is delayed because the database schema has not been approved. Severity: high.',
 'Dependency gap: frontend integration depends on the API contract that is pending review.',
 'Delivery challenge: integration testing cannot start until staging server access is restored.',
 'Delivery forecast: the pilot date is unconfirmed; no revised delivery date is confirmed.',
 'Pending decision: the product owner must approve the ticket retention period.',
 'Unresolved issue: the staging server is unavailable.',
 'Action item: Ravi will restore staging server access by 18 September 2026. Status: pending.',
 'Action item: arrange the retention-policy review. Owner: unknown. Due date: unknown.',
 'Action item: test the webhook retry path. Owner: Meera; Due date: 2026-09-23; Status: pending.',
]

def create(output):
 output=Path(output);output.mkdir(parents=True,exist_ok=True)
 (output/'agent_project.txt').write_text('Atlas Portal - controlled Milestone 2 validation\n\n'+'\n\n'.join(STATEMENTS),encoding='utf-8')
 doc=Document();doc.add_heading('Atlas Portal - controlled agent validation',0)
 for line in STATEMENTS:doc.add_paragraph(line)
 doc.save(output/'agent_project.docx')
 with (output/'agent_project.csv').open('w',newline='',encoding='utf-8') as f:
  w=csv.writer(f);w.writerow(['Statement']);w.writerows([[s] for s in STATEMENTS])
 styles=getSampleStyleSheet();story=[Paragraph('Atlas Portal - controlled agent validation',styles['Title']),Spacer(1,15)]
 for line in STATEMENTS:story.extend([Paragraph(escape(line),styles['BodyText']),Spacer(1,9)])
 SimpleDocTemplate(str(output/'agent_project.pdf'),pagesize=A4).build(story)
 return sorted(output.iterdir())

if __name__=='__main__':
 for p in create(Path(__file__).resolve().parents[1]/'validation_samples'):print(p)
