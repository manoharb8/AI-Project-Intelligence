"""
One-off script to generate realistic, internally-consistent sample project
documents for the AI-Driven Enterprise Project Intelligence & Risk
Management Platform demo. Not part of the shipped application - run once
to populate data/sample_documents/.
"""

import csv
from pathlib import Path

from docx import Document
from fpdf import FPDF
from fpdf.enums import XPos, YPos

OUT_DIR = Path(__file__).resolve().parent.parent / "data" / "sample_documents"
OUT_DIR.mkdir(parents=True, exist_ok=True)


def make_pdf(filename: str, title: str, sections: list[tuple[str, str]]) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, title, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(2)
    for heading, body in sections:
        pdf.set_font("Helvetica", "B", 12)
        pdf.multi_cell(0, 8, heading, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 6, body, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(3)
    pdf.output(str(OUT_DIR / filename))
    print(f"wrote {filename}")


def make_docx(filename: str, title: str, sections: list[tuple[str, str]]) -> None:
    doc = Document()
    doc.add_heading(title, level=1)
    for heading, body in sections:
        doc.add_heading(heading, level=2)
        for line in body.split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())
    doc.save(OUT_DIR / filename)
    print(f"wrote {filename}")


# ---------------------------------------------------------------------------
# 1. project_proposal.pdf
# ---------------------------------------------------------------------------
make_pdf(
    "project_proposal.pdf",
    "Project Proposal: AI-Driven Enterprise Project Intelligence "
    "& Risk Management Platform",
    [
        (
            "Prepared For",
            "Orion Systems Pvt Ltd - Internal PMO Pilot Program. "
            "Prepared by the Project Intelligence Engineering Team.",
        ),
        (
            "Objective",
            "Build a platform that ingests a project's own documents "
            "(proposals, requirements, sprint updates, task lists, risk "
            "reports, meeting notes) and uses Retrieval-Augmented "
            "Generation to answer questions about project scope, "
            "deliverables, risks, blockers and progress, instead of "
            "relying on a project manager to manually search through "
            "scattered files.",
        ),
        (
            "Scope - Milestone 1",
            "Document ingestion for PDF, DOCX, CSV and TXT files; text "
            "normalization; chunking; embedding generation; vector "
            "indexing with Chroma; and a working retrieval interface that "
            "returns the most relevant chunks for a natural-language "
            "query, with source and relevance information. Milestone 1 "
            "does not include automated risk scoring, forecasting, or a "
            "conversational assistant - those are explicitly deferred.",
        ),
        (
            "Scope - Future Milestones",
            "Milestone 2 will add automated risk detection, blocker "
            "identification and a project health dashboard on top of the "
            "Milestone 1 retrieval layer. Milestone 3 will add a "
            "conversational project-intelligence assistant and delivery "
            "forecasting.",
        ),
        (
            "Key Deliverables for Milestone 1",
            "D1: Document ingestion module supporting PDF, DOCX, CSV and "
            "TXT. D2: A working RAG pipeline (extraction, normalization, "
            "chunking, embedding, indexing, retrieval). D3: A retrieval "
            "testing report covering all supported file types and an "
            "insufficient-information case. D4: A Streamlit application "
            "demonstrating upload, processing and query end to end.",
        ),
        (
            "Timeline",
            "The Milestone 1 pilot is planned across 6 two-week sprints. "
            "The project is currently in Sprint 3 of 6.",
        ),
    ],
)

# ---------------------------------------------------------------------------
# 2. project_requirements.docx
# ---------------------------------------------------------------------------
make_docx(
    "project_requirements.docx",
    "Software Requirements Specification",
    [
        (
            "Functional Requirements",
            "FR1: The system shall allow a user to upload project "
            "documents in PDF, DOCX, CSV and TXT formats.\n"
            "FR2: The system shall extract text content from each "
            "uploaded document.\n"
            "FR3: The system shall normalize extracted text to remove "
            "formatting noise before further processing.\n"
            "FR4: The system shall split normalized text into "
            "overlapping chunks suitable for embedding.\n"
            "FR5: The system shall generate a numeric embedding for "
            "every chunk.\n"
            "FR6: The system shall store chunk embeddings in a "
            "persistent vector index.\n"
            "FR7: The system shall accept a natural-language query and "
            "return the most relevant chunks, ranked by similarity, "
            "along with the source filename and chunk position.\n"
            "FR8: The system shall clearly indicate when no uploaded "
            "document contains sufficient information to answer a "
            "query, rather than returning an unrelated chunk as if it "
            "were an answer.",
        ),
        (
            "Non-Functional Requirements",
            "NFR1: The embedding and retrieval pipeline shall run "
            "entirely on local infrastructure for Milestone 1, without "
            "requiring a paid third-party API.\n"
            "NFR2: The user interface shall present processing status "
            "per uploaded file (extracted, chunked, embedded, "
            "indexed).\n"
            "NFR3: The system shall be demonstrable on a single "
            "developer laptop without external infrastructure.\n"
            "NFR4: Source code shall be organized so the RAG pipeline "
            "can be reused behind a different user interface in a "
            "later milestone.",
        ),
        (
            "Out of Scope for Milestone 1",
            "Automated project health scoring, delivery date "
            "forecasting, a conversational chat assistant, and a "
            "multi-project dashboard are explicitly out of scope for "
            "Milestone 1 and are planned for later milestones.",
        ),
    ],
)

# ---------------------------------------------------------------------------
# 3. sprint_update.txt
# ---------------------------------------------------------------------------
sprint_update = """Sprint 3 Update - AI-Driven Enterprise Project Intelligence Platform

Sprint Goal: Deliver a working RAG pipeline (ingestion through retrieval)
and begin the Streamlit user interface.

Completed this sprint:
- Document ingestion module supporting PDF, DOCX, CSV and TXT files.
- Text normalization step to clean extracted content before chunking.
- Chunking module with configurable chunk size and overlap.
- Local embedding generation, verified to run without any external
  network dependency.
- Chroma vector store integration for persistent chunk indexing.

In progress:
- Streamlit interface for document upload and processing status.
- Streamlit interface for querying the knowledge base and displaying
  retrieved chunks with source and relevance information.

Blockers:
- No active blockers this sprint.

Risks carried into next sprint:
- The relevance threshold used to detect insufficient-information
  queries needs tuning against real sample documents before the
  Milestone 1 demo, to avoid either hallucinated confidence or overly
  cautious "no answer" responses.

Next sprint plan:
- Finish the Streamlit upload, processing and query pages.
- Write the Milestone 1 unit test suite covering every pipeline stage.
- Prepare the Milestone 1 demonstration and documentation.
"""
(OUT_DIR / "sprint_update.txt").write_text(sprint_update, encoding="utf-8")
print("wrote sprint_update.txt")

# ---------------------------------------------------------------------------
# 4. task_list.csv
# ---------------------------------------------------------------------------
tasks = [
    ["Task ID", "Task Name", "Owner", "Status", "Priority", "Due Date"],
    ["T-01", "Build PDF/DOCX/CSV/TXT ingestion module", "Manohar", "Done", "High", "2026-01-16"],
    ["T-02", "Implement text normalization", "Manohar", "Done", "High", "2026-01-18"],
    ["T-03", "Implement chunking module", "Manohar", "Done", "High", "2026-01-20"],
    ["T-04", "Implement local embedding generation", "Manohar", "Done", "High", "2026-01-23"],
    ["T-05", "Implement Chroma vector store indexing", "Manohar", "Done", "High", "2026-01-25"],
    ["T-06", "Implement retrieval with relevance threshold", "Manohar", "In Progress", "High", "2026-01-30"],
    ["T-07", "Build Streamlit upload and processing pages", "Manohar", "In Progress", "Medium", "2026-02-02"],
    ["T-08", "Build Streamlit query and retrieval results page", "Manohar", "Not Started", "Medium", "2026-02-05"],
    ["T-09", "Write Milestone 1 unit tests", "Manohar", "Not Started", "High", "2026-02-07"],
    ["T-10", "Write README and architecture documentation", "Manohar", "Not Started", "Medium", "2026-02-08"],
    ["T-11", "Prepare Milestone 1 mentor demonstration", "Manohar", "Not Started", "High", "2026-02-10"],
]
with open(OUT_DIR / "task_list.csv", "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerows(tasks)
print("wrote task_list.csv")

# ---------------------------------------------------------------------------
# 5. project_risk_report.pdf
# ---------------------------------------------------------------------------
make_pdf(
    "project_risk_report.pdf",
    "Project Risk Report - Sprint 3",
    [
        (
            "Overall Risk Level",
            "The overall project risk level for Milestone 1 is currently "
            "assessed as MEDIUM. No risk is currently rated as blocking "
            "delivery, but two risks require active mitigation before "
            "the Milestone 1 demonstration.",
        ),
        (
            "Risk R1: Embedding Dependency Reliability",
            "Description: Some embedding approaches depend on "
            "downloading a pretrained model from an external model hub "
            "on first run, which can fail on machines with restricted "
            "or slow internet access. Likelihood: Medium. Impact: "
            "Medium. Mitigation: Milestone 1 uses a local, dependency-"
            "light embedding method that requires no external download, "
            "removing this risk for the current milestone.",
        ),
        (
            "Risk R2: Chunking Losing Context Across Boundaries",
            "Description: Splitting long documents into chunks can "
            "separate a risk description from its mitigation, or a task "
            "from its deadline, if chunk boundaries fall in the wrong "
            "place. Likelihood: Medium. Impact: High. Mitigation: "
            "Chunking uses overlapping word windows so context near "
            "chunk boundaries is preserved in at least one chunk.",
        ),
        (
            "Risk R3: Hallucinated Answers on Out-of-Scope Queries",
            "Description: If the retrieval step always returns the "
            "closest chunks even when none of them are actually "
            "relevant, the platform could imply an answer exists when "
            "the uploaded documents do not contain one. Likelihood: "
            "Medium. Impact: High. Mitigation: Retrieved chunks are "
            "compared against a relevance distance threshold, and the "
            "interface explicitly states when no uploaded document "
            "contains sufficient information to answer the query.",
        ),
        (
            "Risk R4: Milestone 1 Scope Creep",
            "Description: Requirements gathering surfaced several "
            "future modules (risk scoring, forecasting, dashboards, "
            "conversational assistant). Building these before the core "
            "RAG pipeline is proven could delay Milestone 1. "
            "Likelihood: Low. Impact: High. Mitigation: Milestone 1 is "
            "scoped strictly to ingestion, chunking, embedding, "
            "indexing and retrieval; all other modules are explicitly "
            "deferred to later milestones.",
        ),
        (
            "Current Blockers",
            "There are no blocking issues as of Sprint 3. The relevance "
            "threshold for insufficient-information detection is still "
            "being tuned and is tracked as an open risk item (R3), not "
            "a blocker.",
        ),
    ],
)

# ---------------------------------------------------------------------------
# 6. meeting_notes.docx
# ---------------------------------------------------------------------------
make_docx(
    "meeting_notes.docx",
    "Sprint 3 Review Meeting Notes",
    [
        (
            "Meeting Details",
            "Date: Sprint 3, Week 2\n"
            "Attendees: Manohar (Developer), Project Mentor\n"
            "Purpose: Review Sprint 3 progress and confirm Milestone 1 "
            "scope ahead of the demonstration.",
        ),
        (
            "Progress Reviewed",
            "The document ingestion module, normalization, chunking, "
            "embedding generation and Chroma vector indexing are all "
            "complete and individually tested. The Streamlit interface "
            "for upload and processing is in progress.",
        ),
        (
            "Decisions",
            "Decision 1: The conversational project-intelligence "
            "assistant and the automated risk-scoring dashboard are "
            "confirmed as out of scope for Milestone 1 and deferred to "
            "Milestone 3 and Milestone 2 respectively.\n"
            "Decision 2: The relevance distance threshold used to "
            "detect insufficient-information queries will be tuned "
            "against the bundled sample project documents before the "
            "demonstration, rather than left at a default value.",
        ),
        (
            "Blockers Discussed",
            "No active blockers were raised in this meeting. The "
            "embedding dependency risk (R1) discussed in the project "
            "risk report is considered resolved for Milestone 1 by "
            "using a local, dependency-light embedding method.",
        ),
        (
            "Action Items",
            "Action 1: Finish the Streamlit query and retrieval results "
            "page - Owner: Manohar - Due before Milestone 1 demo.\n"
            "Action 2: Write the Milestone 1 unit test suite covering "
            "every pipeline stage - Owner: Manohar - Due before "
            "Milestone 1 demo.\n"
            "Action 3: Write the README and architecture documentation "
            "- Owner: Manohar - Due before Milestone 1 demo.",
        ),
    ],
)

print("Done generating sample documents.")
