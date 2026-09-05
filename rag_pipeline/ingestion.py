"""
Document ingestion.

Responsible for extracting raw text from the four supported project
artifact types: PDF, DOCX, CSV and TXT. Every extractor accepts either a
filesystem path (str) or a file-like/binary buffer (e.g. the object handed
over by Streamlit's file uploader), so the same functions work both from
disk (used by the test suite and sample data) and from the UI.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Union

import pandas as pd
from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".csv", ".txt"}

FileInput = Union[str, "io.IOBase"]


class UnsupportedFileTypeError(ValueError):
    """Raised when a file extension is not one of the supported types."""


def get_file_type(filename: str) -> str:
    """Return the normalized file type ('pdf', 'docx', 'csv', 'txt') for a filename."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext}'. Supported types: {sorted(SUPPORTED_EXTENSIONS)}"
        )
    return ext.lstrip(".")


def extract_text(file_input: FileInput, filename: str) -> str:
    """
    Extract raw text from a supported project document.

    Args:
        file_input: a filesystem path, or a file-like object (must support
            .read() / be seekable, as provided by Streamlit's uploader).
        filename: the original filename, used to determine file type.

    Returns:
        The extracted raw text (not yet normalized).

    Raises:
        UnsupportedFileTypeError: if the file extension is not supported.
    """
    file_type = get_file_type(filename)
    extractor = {
        "pdf": _extract_pdf,
        "docx": _extract_docx,
        "csv": _extract_csv,
        "txt": _extract_txt,
    }[file_type]
    return extractor(file_input)


def _extract_pdf(file_input: FileInput) -> str:
    reader = PdfReader(file_input)
    pages_text = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages_text)


def _extract_docx(file_input: FileInput) -> str:
    document = DocxDocument(file_input)
    parts = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cells)
            if row_text.strip(" |"):
                parts.append(row_text)

    return "\n".join(parts)


def _extract_csv(file_input: FileInput) -> str:
    """
    Extract CSV content as text using pandas, rendering each row as
    "column: value" pairs so column names stay attached to their values
    once the text is later split into chunks (a chunk boundary can land
    anywhere, so context should not depend on a nearby header row).
    """
    buffer, should_close = _as_text_buffer(file_input)
    try:
        dataframe = pd.read_csv(buffer, dtype=str, keep_default_na=False)
    finally:
        if should_close:
            buffer.close()

    if dataframe.empty:
        return ""

    lines = []
    for _, row in dataframe.iterrows():
        if not any(str(value).strip() for value in row):
            continue
        pairs = [f"{col.strip()}: {str(value).strip()}" for col, value in row.items()]
        lines.append(", ".join(pairs))
    return "\n".join(lines)


def _extract_txt(file_input: FileInput) -> str:
    buffer, should_close = _as_text_buffer(file_input)
    try:
        return buffer.read()
    finally:
        if should_close:
            buffer.close()


def _as_text_buffer(file_input: FileInput):
    """
    Normalize a path or file-like object into a readable text buffer.

    Returns a tuple of (buffer, should_close) so callers know whether they
    opened a new file handle that they are responsible for closing.
    """
    if hasattr(file_input, "read"):
        content = file_input.read()
        if hasattr(file_input, "seek"):
            file_input.seek(0)
        if isinstance(content, bytes):
            content = content.decode("utf-8", errors="replace")
        return io.StringIO(content), True
    return open(file_input, "r", encoding="utf-8", errors="replace"), True
